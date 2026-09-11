"""
AI PRODUCT FACTORY — COMMERCIAL SPREADSHEET BUILDER
Modality: Spreadsheets, Financial Models, Workflow Trackers, Directories
Generates professional .xlsx workbooks with styled headers, frozen panes,
and auto-fit columns (or typed CSV fallback).
"""

import os
import sys
import csv
import json
import argparse

# Ensure local system vendor packages (openpyxl, et_xmlfile) are discoverable
vendor_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "vendor"))
if os.path.exists(vendor_dir) and vendor_dir not in sys.path:
    sys.path.insert(0, vendor_dir)

def generate_spreadsheet(output_path, title="Product Model", sheets_data=None):
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    
    # Default sample data if none provided
    if not sheets_data:
        sheets_data = {
            "Model Overview": [
                ["AI Product Factory — Professional Model", ""],
                ["Title", title],
                ["Status", "VERIFIED"],
                ["", ""],
                ["Category", "Baseline", "Optimized", "Delta (%)", "Notes"],
                ["Operating Overhead ($)", 12500, 3200, -74.4, "Automated pipeline savings"],
                ["Weekly Hours Drag (hrs)", 38, 6, -84.2, "Zero manual reporting"],
                ["Compliance Accuracy Score", 82, 99.5, 21.3, "Continuous automated audit"],
                ["Customer Acquisition Cost ($)", 450, 180, -60.0, "Creator-driven channel"]
            ]
        }

    # Attempt openpyxl compilation
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default blank sheet

        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=10)
        title_font = Font(size=14, bold=True, color="0F172A")
        bold_font = Font(bold=True)

        for sheet_name, rows in sheets_data.items():
            ws = wb.create_sheet(title=sheet_name)
            ws.views.sheetView[0].showGridLines = True

            for r_idx, row in enumerate(rows, start=1):
                for c_idx, val in enumerate(row, start=1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=val)
                    # Title row
                    if r_idx == 1 and c_idx == 1:
                        cell.font = title_font
                    # Table headers
                    elif len(row) > 2 and r_idx == 5:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="center" if c_idx > 1 else "left")
                    # Numbers formatting
                    elif isinstance(val, (int, float)):
                        if isinstance(val, float) and abs(val) < 1.0 and val != 0:
                            cell.number_format = "0.0%"
                        elif isinstance(val, (int, float)) and abs(val) > 100:
                            cell.number_format = "$#,##0" if "($)" in str(rows[4][c_idx-1]) else "#,##0"

            # Auto-fit column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            # Freeze panes below header
            if len(rows) >= 5:
                ws.freeze_panes = "A6"

        wb.save(output_path)
        print(f"[SPREADSHEET BUILDER] Generated professional XLSX: {output_path}")
        return True

    except ImportError:
        # Fallback to clean CSV
        csv_path = os.path.splitext(output_path)[0] + ".csv"
        first_sheet = list(sheets_data.values())[0]
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            for row in first_sheet:
                writer.writerow(row)
        print(f"[SPREADSHEET BUILDER] openpyxl unavailable. Generated structured CSV: {csv_path}")
        return True

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="AI Product Factory Spreadsheet Builder")
    parser.add_argument("--output", required=True, help="Output file path (.xlsx or .csv)")
    parser.add_argument("--title", default="Product Model", help="Model title")
    args = parser.parse_args()

    generate_spreadsheet(args.output, title=args.title)
