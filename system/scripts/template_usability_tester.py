"""
AI PRODUCT FACTORY — SPREADSHEET & TEMPLATE USABILITY TESTER
Automated formula auditing, error token scanning (#REF!, #DIV/0!), frozen panes verification,
and commercial usability compliance.
"""

import os
import sys
import re
import argparse
import json

# Ensure vendor directory is in sys.path
vendor_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "vendor"))
if os.path.exists(vendor_dir) and vendor_dir not in sys.path:
    sys.path.insert(0, vendor_dir)

ERROR_TOKENS = ["#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!"]

def audit_spreadsheet(file_path):
    """
    Audits a spreadsheet file (.xlsx or .csv) for commercial publication readiness.
    Returns a dict with overall status, defect list, and metrics.
    """
    if not os.path.exists(file_path):
        return {"status": "FAIL", "error": f"File not found: {file_path}"}

    ext = os.path.splitext(file_path)[1].lower()
    defects = []
    warnings = []
    metrics = {"sheets": 0, "cells_audited": 0, "formulas_checked": 0}

    if ext == ".xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=False)
            wb_values = openpyxl.load_workbook(file_path, data_only=True)
            metrics["sheets"] = len(wb.sheetnames)

            # Check 1: Sheet Naming
            for sname in wb.sheetnames:
                if re.match(r"^Sheet\d+$", sname, re.IGNORECASE):
                    warnings.append(f"Generic sheet name detected: '{sname}'. Use purposeful names.")

            # Check 2: Instructions/Assumptions Tab
            has_guide = any(k in " ".join(wb.sheetnames).lower() for k in ["instruction", "overview", "assumption", "guide", "readme"])
            if not has_guide:
                warnings.append("Missing explicit 'Instructions' or 'Assumptions' sheet to orient customer.")

            # Check 3: Formulas, Errors, and Formatting across all sheets
            for sheet_name in wb.sheetnames:
                ws_formulas = wb[sheet_name]
                ws_vals = wb_values[sheet_name]

                # Check Freeze Panes
                if ws_formulas.max_row > 10 and not ws_formulas.freeze_panes:
                    warnings.append(f"Sheet '{sheet_name}' has {ws_formulas.max_row} rows but NO frozen panes configured.")

                for row_f, row_v in zip(ws_formulas.iter_rows(), ws_vals.iter_rows()):
                    for cell_f, cell_v in zip(row_f, row_v):
                        metrics["cells_audited"] += 1
                        val_str = str(cell_v.value or "")
                        formula_str = str(cell_f.value or "")

                        # Formula check
                        if formula_str.startswith("="):
                            metrics["formulas_checked"] += 1

                        # Error token check
                        for err in ERROR_TOKENS:
                            if err in val_str or err in formula_str:
                                defects.append(f"CRITICAL: Formula error token '{err}' found at {sheet_name}!{cell_f.coordinate} (Formula: {formula_str})")

        except Exception as e:
            defects.append(f"CRITICAL: Failed to parse XLSX workbook: {e}")

    elif ext == ".csv":
        import csv
        metrics["sheets"] = 1
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            for r_idx, row in enumerate(reader, start=1):
                for c_idx, cell in enumerate(row, start=1):
                    metrics["cells_audited"] += 1
                    for err in ERROR_TOKENS:
                        if err in cell:
                            defects.append(f"CRITICAL: Error token '{err}' found at Row {r_idx}, Col {c_idx}")

    else:
        return {"status": "FAIL", "error": f"Unsupported format: {ext}"}

    # Final score
    passed = len(defects) == 0
    return {
        "status": "PASS" if passed else "FAIL",
        "file": file_path,
        "critical_defects": defects,
        "warnings": warnings,
        "metrics": metrics
    }

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="AI Product Factory Spreadsheet Usability Tester")
    parser.add_argument("--input", required=True, help="Path to spreadsheet file (.xlsx or .csv)")
    parser.add_argument("--json", action="store_true", help="Output raw JSON results")
    args = parser.parse_args()

    result = audit_spreadsheet(args.input)
    if args.json:
        print(json.dumps(result, indent=2))
    else:
        print(f"\n=======================================================")
        print(f"SPREADSHEET USABILITY & AUDIT REPORT: {result['file']}")
        print(f"Status: {result['status']}")
        print(f"Metrics: Sheets={result['metrics']['sheets']}, Cells={result['metrics']['cells_audited']}, Formulas={result['metrics']['formulas_checked']}")
        print(f"=======================================================")
        if result['critical_defects']:
            print(f"\nCRITICAL DEFECTS ({len(result['critical_defects'])}):")
            for d in result['critical_defects']:
                print(f"  [X] {d}")
        if result['warnings']:
            print(f"\nUSABILITY WARNINGS ({len(result['warnings'])}):")
            for w in result['warnings']:
                print(f"  [!] {w}")
        if not result['critical_defects'] and not result['warnings']:
            print("\n[OK] Zero formula defects. Professional formatting standards satisfied.")

    sys.exit(0 if result['status'] == "PASS" else 1)
